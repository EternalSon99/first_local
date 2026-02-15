from pathlib import Path


def parse_pdf(file_path: Path) -> list[dict]:
    """Parse a PDF file and return pages with text and page numbers.

    Returns list of {"page_number": int, "text": str}.
    """
    import fitz  # PyMuPDF

    pages = []
    with fitz.open(str(file_path)) as doc:
        for page_num, page in enumerate(doc, start=1):
            text = page.get_text().strip()
            if text:
                pages.append({"page_number": page_num, "text": text})
    return pages


def parse_docx(file_path: Path) -> list[dict]:
    """Parse a DOCX file and return content as a single page.

    Returns list of {"page_number": int, "text": str}.
    """
    from docx import Document

    doc = Document(str(file_path))
    full_text = "\n".join(para.text for para in doc.paragraphs if para.text.strip())
    if not full_text.strip():
        return []
    return [{"page_number": 1, "text": full_text}]


def parse_pptx(file_path: Path) -> list[dict]:
    """Parse a PPTX file with each slide as a page.

    Returns list of {"page_number": int, "text": str}.
    """
    from pptx import Presentation

    prs = Presentation(str(file_path))
    slides = []
    for slide_num, slide in enumerate(prs.slides, start=1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        texts.append(text)
        if texts:
            slides.append({"page_number": slide_num, "text": "\n".join(texts)})
    return slides


def parse_txt(file_path: Path) -> list[dict]:
    """Parse a plain text file.

    Returns list of {"page_number": int, "text": str}.
    """
    text = file_path.read_text(encoding="utf-8").strip()
    if not text:
        return []
    return [{"page_number": 1, "text": text}]


PARSERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".pptx": parse_pptx,
    ".txt": parse_txt,
}


def parse_file(file_path: Path) -> list[dict]:
    """Parse a file based on its extension.

    Returns list of {"page_number": int, "text": str}.
    Raises ValueError if file type is unsupported.
    """
    ext = file_path.suffix.lower()
    parser = PARSERS.get(ext)
    if parser is None:
        raise ValueError(f"Unsupported file type: {ext}")
    return parser(file_path)
